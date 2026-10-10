"""Builds deliverables/attrition_narrative_v9.docx (v8 + Quinn's wording-only fixes V1, V2 and optional items;
research from research/exec_research_tiein_v2.md; v1-v8 untouched).
Numbers ONLY from analysis/exec2/exec_findings_v2_draft.md (QA-cleared by Quinn, Oct 9, 2026)."""
import os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn

OUT = "../attrition_narrative_v9.docx"
if os.path.exists(OUT) and "--force" not in sys.argv:
    sys.exit(f"{OUT} exists; refusing to overwrite")

ACC = RGBColor(0xC2, 0x41, 0x0C); TXT = RGBColor(0x40, 0x40, 0x40)
MUT = RGBColor(0x6B, 0x6B, 0x6B); DARK = RGBColor(0x26, 0x26, 0x26)
doc = Document(); s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.left_margin = s.right_margin = Inches(0.7); s.top_margin = Inches(0.45); s.bottom_margin = Inches(0.45)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10); st.font.color.rgb = TXT
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.space_before = Pt(0); st.paragraph_format.line_spacing = 1.0
W = 7.1

def P(parts, size=10, color=None, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    if isinstance(parts, str): parts = [(parts, "")]
    for t, f in parts:
        r = p.add_run(t); r.font.size = Pt(size)
        r.bold = "b" in f; r.font.color.rgb = ACC if "a" in f else (DARK if "d" in f else (MUT if "m" in f else (color or TXT)))
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p

def H(text, before=7):
    P([(text, "bd")], size=13, before=before, after=3, keep=True)

def img(path, width, alt, before=2, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    shape = p.add_run().add_picture(path, width=Inches(width))
    docpr = shape._inline.docPr; docpr.set("descr", alt); docpr.set("title", alt.split(":")[0])

from docx.oxml import OxmlElement
ACC2 = RGBColor(0x7C, 0x2D, 0x12)

def shaded_box(parts_list, fill="F2F2F2"):
    """Gray box; a paragraph starting 'Look into' begins a second row, and rows can't split across pages,
    so that heading is never orphaned from its list."""
    groups = [[]]
    for item in parts_list:
        if item[0][0][0].startswith("Look into"):
            groups.append([])
        groups[-1].append(item)
    t = doc.add_table(rows=len(groups), cols=1); t.autofit = False
    tb = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:val"), "nil"); b.append(e)
    tb.append(b)
    for ri, grp in enumerate(groups):
        row = t.rows[ri]
        trpr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trpr.append(cs)
        c = row.cells[0]; c.width = Inches(W)
        pr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); pr.append(shd)
        mar = OxmlElement("w:tcMar")
        top = 110 if ri == 0 else 40; bot = 90 if ri == len(groups) - 1 else 40
        for k, v in (("top", top), ("bottom", bot), ("start", 160), ("end", 160)):
            e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
        pr.append(mar)
        first = True
        for item in grp:
            parts, size, after = item[:3]; indent = item[3] if len(item) > 3 else 0
            p = c.paragraphs[0] if first else c.add_paragraph(); first = False
            if indent:
                p.paragraph_format.left_indent = Inches(indent); p.paragraph_format.first_line_indent = Inches(-indent)
            for txt, f in parts:
                r = p.add_run(txt); r.font.size = Pt(size); r.bold = "b" in f
                r.font.color.rgb = DARK if "d" in f else MUT
            p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def rec_list(items, lead_color="bd"):
    """Compact: bold lead-in + one sentence + citations in a trailing parenthesis (no separate 'Why:' line)."""
    for i, (lead, body, cites) in enumerate(items, 1):
        p = P([(f"{i}.  ", "bd"), (lead, lead_color), (body + " ", ""), (f"({cites})", "m")], size=9.5, after=1)
        p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.first_line_indent = Inches(-0.22)

# Title
P("IBM HR teaching data  \u00b7  fictional, simulated employees  \u00b7  executive summary", 9, MUT, after=3)
P([("Overtime workers", "ba"), (" leave most in IBM\u2019s fictional HR data, lower-paid junior ones most of all", "bd")],
  size=20, after=10)

# Executive summary (Quinn's required sentence verbatim as sentence 3)
P([("EXECUTIVE SUMMARY", "bd")], size=9.5, after=2)
P("In IBM\u2019s fictional teaching dataset of 1,470 simulated employees, 16% left. Overtime workers are 28% of staff "
  "but 54% of everyone who left, and within that group lower-paid, mostly junior staff leave most: about 55% of them "
  "left, close to 30% of all leavers. "
  "Newer staff also leave more often, with or without overtime; we flag that as one career-stage picture to watch, not a cleared finding. "
  "As an illustration, not a forecast, if overtime workers left at the company average, attrition would fall from "
  "about 16% to about 12%, roughly 60 fewer leavers. The first step is to find out where overtime is concentrated "
  "and why, then pilot changes in teams where overtime is heavy and not chosen; these are simulated people, and the "
  "data show who left, not why.", after=6)

# What the Data Shows
H("What the Data Shows: overtime workers leave most, lower-paid junior ones most of all")
P("In this fictional dataset, 416 people, 28% of staff, worked overtime. 31% of them left, about 1.9 times the "
  "company rate of 16% and about 3 times the rate of staff without overtime (10%). Together they account for 54% of "
  "everyone who left.")
P("Within overtime, lower-paid, mostly junior staff leave most. Overtime workers earning under roughly $3,000 to "
  "$3,500 a month (about 115 to 130 people, 8% to 9% of staff) left at about 55%, about 2.7 to 2.9 times the rate of "
  "the other overtime workers (about 20%). They account for close to 30% of everyone who left. They are part of the "
  "overtime finding, not a separate pattern.")
img("leaving_chart_v9.png", 6.6, "Chart: in IBM's fictional data, 16% of all staff and 10% of staff without overtime "
    "left, against 31% of the 416 overtime workers and about 55% of the about 115 to 130 lower-paid, junior overtime "
    "workers. Overtime workers are 54% of everyone who left; the lower-paid, junior ones, inside that group, are close "
    "to 30%.", before=4, after=6)
P([("Context: one career-stage picture. ", "bd"),
   ("Pay, job level, tenure and total experience move together, and the data cannot say which of them matters. "
    "Almost all of these lower-paid overtime workers (97%) are in the most junior job level, so pay and job level "
    "cannot be separated: overtime workers in the most junior job level (156 people) left at 53%, against 17% for "
    "the 260 overtime workers at higher job levels, close to the company average of 16%.", "")],
  size=9.5, color=MUT)

# Worth watching box (item 3, gray only)
shaded_box([
    ([("Worth watching, not a cleared finding: newer staff", "bd")], 10, 3),
    ([("Staff with 1 year or less at the company (215 people, 15% of staff) left at 35%, about 40 people above the "
       "company rate; about a third of them (32%) work overtime. The group is large, but the exact group isn\u2019t "
       "stable across repeated analyses, because new hires, short tenure, junior level and low pay overlap. Treat them "
       "as one career-stage picture to watch, not a proven cause; it is not added to the overtime figures. Among staff "
       "without overtime, newer staff also leave more often: a real difference inside that group, but modest against "
       "the company rate.", "")], 9.5, 4),
    ([("Research finds that longer time with a company goes with less leaving (Rubenstein et al., 2018); people newer "
       "to a company may have had less chance to build the ties to colleagues, fit with the job and what people would "
       "give up by leaving that go with staying (Mitchell et al., 2001). Research on people in their first year or so "
       "finds that a structured start and finding the information they need go with a clearer role and with feeling "
       "accepted by colleagues, and that settling in goes with intending to stay and with staying, most clearly where "
       "people feel accepted by colleagues (Bauer et al., 2007).", "")], 9.5, 4),
    ([("Look into (watch item, not a cleared finding), ", "bd"), ("for all new hires, never by personal characteristics", "")], 9.5, 2),
    ([("1.  ", "bd"), ("Review onboarding for the first year across roles and teams. ", "bd"),
      ("Check whether new hires get a clear role, the information they need and a chance to meet colleagues "
       "(Bauer et al., 2007).", "")], 9.5, 2, 0.2),
    ([("2.  ", "bd"), ("Hold a 90-day and a 6-month check-in for all new hires. ", "bd"),
      ("The same offer to everyone who joins (Mitchell et al., 2001; Bauer et al., 2007).", "")], 9.5, 2, 0.2),
    ([("3.  ", "bd"), ("Track first-year leaving by team. ", "bd"),
      ("Report it next to overtime levels, so leaders can see whether it is a company-wide pattern or sits in a few "
       "teams (Rubenstein et al., 2018).", "")], 9.5, 0, 0.2),
])

# Why It Matters and What to Do
H("Why It Matters and What to Do: overtime is the clearest place to start")
P("As an illustration, not a forecast, if overtime workers left at the company average, overall attrition would fall "
  "from about 16% to about 12%, roughly 60 fewer leavers. On the same illustration, the lower-paid, mostly junior "
  "overtime workers would take it from about 16% to about 13%, roughly 45 to 50 fewer leavers. They are already "
  "inside the overtime group, so the two illustrations overlap and are not added.")
P([("These are ideas to test, not proven fixes. ", "bd"),
   ("Each applies to teams and groups as a whole, never to individuals and never selected by age, gender or "
    "marital status. Sources in brackets give research background for each idea; none of them tested the idea "
    "itself.", "")], after=5)

P([("Overtime workers (cleared finding)", "bd")], size=11, after=2, keep=True)
P("Long hours can drain time and energy from family life (Greenhaus & Beutell, 1985), and that kind of conflict is "
  "linked to leaving (Rubenstein et al., 2018). But research measures workload, not overtime, and finds it only "
  "weakly related to leaving, if anything slightly protective (Rubenstein et al., 2018). Obstacles like red tape and "
  "unclear roles go with higher turnover, while stretching demands do not (Podsakoff et al., 2007). Research on workload would "
  "not lead us to expect a gap this large, so treat this fictional result as a pointer to where to look, not proof "
  "of cause or evidence about any real employer.", after=4)
rec_list([
    ("Review workload and staffing where overtime is concentrated. ",
     "Map overtime by team and role, and ask whether it is a predictable crunch or a chronic pattern, and whether it "
     "is chosen or required.",
     "Podsakoff et al., 2007"),
    ("Offer manager check-ins (\u201cstay conversations\u201d) to everyone in the overtime group. ",
     "Ask what keeps people and what might make them leave, then act on what is in the team\u2019s control.",
     "Mitchell et al., 2001; Hausknecht et al., 2009"),
    ("Pick one measure and pilot before rolling out. ",
     "Track the leaving rate of overtime workers by team, try changes in a few teams, and compare with similar teams "
     "that did not change; research finds single factors only modestly related to leaving, so local results decide.",
     "Griffeth et al., 2000; Rubenstein et al., 2018"),
])

P([("Within overtime: lower-paid, mostly junior workers ", "bd"),
   ("(part of the overtime finding; overlaps it, not added)", "")], size=11, before=4, after=2, keep=True)
P("Classic theory holds that people stay while what they get outweighs what they give (March & Simon, 1958), and "
  "that feeling under-rewarded can prompt people to leave (Adams, 1965); this data does not record how people felt "
  "about their pay or their overtime. Research finds that lower pay goes with more leaving, and so does being newer "
  "to a company, though both links are modest (Rubenstein et al., 2018). This is consistent with the pattern, as an "
  "illustration only.", after=3)
P("Aimed at junior-level overtime workers (156 people) as a job-level group, never by pay cut, age, gender or marital status.",
  size=9, color=MUT, after=3)
rec_list([
    ("Check overtime pay, recognition and scheduling at the junior job level. ",
     "Ask whether overtime is paid or recognized, whether notice is given, and whether extra hours are shared fairly "
     "across the team.",
     "March & Simon, 1958; Adams, 1965"),
    ("Hold a defined next-step conversation for junior roles. ",
     "Make clear what the next level involves and how people get there, offered to everyone at that level.",
     "Mitchell et al., 2001"),
    ("Pilot in a few junior-level overtime teams, with measurement. ",
     "Track their leaving rate against similar teams that did not change before rolling out.",
     "Rubenstein et al., 2018"),
])

# What Not to Conclude (as v7)
H("What Not to Conclude: limits of a fictional sample")
nots = [
    ("It is a fictional sample. ", "IBM created these 1,470 employees for teaching. They describe no real workforce "
     "or employer."),
    ("It shows who left, not why. ", "Overtime may simply mark busy or understaffed roles, and pay moves with job "
     "level and experience."),
    ("Pay and career stage overlap. ", "Lower pay, junior job level, short tenure and little experience go together, "
     "and the data can\u2019t say which matters."),
    ("It is not for decisions about individuals. ", "The patterns describe groups. Age, gender and marital status "
     "are descriptive only and never grounds for targeting anyone."),
    ("Small groups are left out. ", "Groups under 100 people were set aside as too small to act on."),
]
for lead, body in nots:
    P([(lead, "bd"), (body, "")], after=2)

# Footer: Ellis's short-form sources (exec_research_tiein_v2.md line 77) + numbers/appendix pointer
P([("Sources: ", "b"), ("Adams (1965); Bauer, Bodner, Erdogan, Truxillo & Tucker (2007); Greenhaus & Beutell (1985); "
   "Griffeth, Hom & Gaertner (2000); Hausknecht, Rodda & Howard (2009); March & Simon (1958); Mitchell, Holtom, Lee, "
   "Sablynski & Erez (2001); Podsakoff, LePine & LePine (2007); Rubenstein, Eberly, Lee & Mitchell (2018). Numbers: "
   "analysis/exec2/exec_findings_v2_draft.md (QA-cleared Oct 9, 2026). Technical detail: "
   "analysis/exec2/technical_appendix.md in the project repo.", "")],
  size=8, color=MUT, before=4, after=0)

cp = doc.core_properties
cp.title = "Overtime workers leave most in IBM\u2019s fictional HR data, lower-paid junior ones most of all"
cp.author = "Iris"; cp.last_modified_by = "Iris"
import datetime as _dt
cp.created = cp.modified = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)
doc.save(OUT); print("saved", OUT)
