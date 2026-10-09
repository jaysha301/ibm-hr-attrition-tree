"""Builds deliverables/attrition_narrative_v4.docx (v3 + Quinn's R1-R7 and O1-O3; v1-v3 untouched).
Numbers ONLY from analysis/exec/exec_findings_draft.md (QA-cleared Oct 9, 2026).
Research and recommendation text is PROVISIONAL (from research/attrition_drivers.md and the
Quinn-cleared v2 narrative) pending Ellis's re-tie to the overtime pattern."""
import os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn

OUT = "../attrition_narrative_v4.docx"
if os.path.exists(OUT) and "--force" not in sys.argv:
    sys.exit(f"{OUT} exists; refusing to overwrite")

ACC = RGBColor(0xC2, 0x41, 0x0C); TXT = RGBColor(0x40, 0x40, 0x40)
MUT = RGBColor(0x6B, 0x6B, 0x6B); DARK = RGBColor(0x26, 0x26, 0x26)
doc = Document(); s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.left_margin = s.right_margin = Inches(0.75); s.top_margin = Inches(0.6); s.bottom_margin = Inches(0.55)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = TXT
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.space_before = Pt(0); st.paragraph_format.line_spacing = 1.1
W = 7.0

def P(parts, size=10.5, color=None, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    if isinstance(parts, str): parts = [(parts, "")]
    for t, f in parts:
        r = p.add_run(t); r.font.size = Pt(size)
        r.bold = "b" in f; r.font.color.rgb = ACC if "a" in f else (DARK if "d" in f else (color or TXT))
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p

def H(text, before=12):
    P([(text, "bd")], size=13, before=before, after=3, keep=True)

def img(path, width, alt, before=2, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    shape = p.add_run().add_picture(path, width=Inches(width))
    docpr = shape._inline.docPr; docpr.set("descr", alt); docpr.set("title", alt.split(":")[0])

# Title
P("IBM HR teaching data  \u00b7  fictional, simulated employees  \u00b7  executive summary", 9, MUT, after=3)
P([("Overtime workers", "ba"), (" leave most: what IBM\u2019s fictional HR data suggests", "bd")], size=21, after=10)

# Executive summary
P([("EXECUTIVE SUMMARY", "bd")], size=9.5, after=2)
P("In IBM\u2019s fictional teaching dataset of 1,470 simulated employees, 16% left, and overtime workers stand out: "
  "they are 28% of staff but 54% of everyone who left. If they left at the company average, overall attrition "
  "would fall from about 16% to about 12%, roughly 60 fewer leavers, as an illustration rather than a forecast. "
  "The first step is to find out where overtime is concentrated and why, then pilot changes in teams where "
  "overtime is heavy and not chosen. The key caution: these are simulated people, and the data show who left, not why.", after=6)
img("numbers_strip_v4.png", W, "Key figures: in IBM's fictional data, 16% of all employees left; 31% of overtime "
    "workers left, against 10% of staff without overtime.", before=4, after=4)

# What the Data Shows
H("What the Data Shows: overtime workers leave at about three times the rate of other staff")
P("In this fictional dataset, 416 people, 28% of staff, worked overtime. 31% of them left, about three times the "
  "rate of staff without overtime (10%). Together they account for 54% of everyone who left. It is the one large "
  "group that held up under every check the team ran.")
img("overtime_chart_v4.png", W, "Chart: in IBM's fictional data, overtime workers are 28% of all staff but 54% "
    "of everyone who left.", before=4, after=6)
P([("Context, not a separate finding: pay and career stage. ", "bd"),
   ("Lower-paid staff (roughly the bottom third, under about $3,000\u2013$4,000 a month) left at about 27%, against "
    "11% for everyone else, wherever the pay line is drawn in that range. But 95% of them are in the most junior job level, so pay and career stage can\u2019t be "
    "separated here. Most of their extra leaving is among those who also work overtime: without overtime, "
    "lower-paid staff leave at about the company average (16%), against 8% for better-paid staff without overtime.", "")],
  size=9.5, color=MUT)

# Why It Matters and What to Do
H("Why It Matters and What to Do: overtime is the clearest place to start")
P("In this fictional data, the overtime group is large enough to move the overall number. If overtime workers "
  "left at the company average, overall attrition would fall from about 16% to about 12%, roughly 60 fewer leavers. That "
  "is an illustration of scale, not a forecast of what any change would achieve.")
P("Research on real organizations gives reason to take overtime seriously, and reason for care. Long hours can "
  "take time and energy from family and other roles, and that kind of conflict is linked to leaving (Greenhaus & "
  "Beutell, 1985; Rubenstein et al., 2018). Equity theory holds that people who feel they give more than they get "
  "back may leave to restore the balance (Adams, 1965); this data does not record how people felt. But research measures workload, not overtime itself: across many studies, workload is "
  "only weakly related to leaving, and slightly in the protective direction (Rubenstein et al., 2018). Demands such "
  "as workload and time pressure are usually classed as \u201cchallenge\u201d demands, which go with lower turnover, "
  "unlike obstacles such as red tape and unclear roles, which go with higher turnover (Podsakoff et al., 2007). "
  "So this fictional data should not be read as confirming the research. This data does not record whether "
  "overtime was chosen or required, or how long it lasted, and in a real company those are the first things to "
  "find out.")
P([("Three suggestions to test, not proven fixes. ", "bd"),
   ("Each is aimed at teams and at overtime workers as a group, never at individuals.", "")], after=4)
recs = [
    ("Find where overtime sits, and why. ",
     "Map overtime by team and role, and ask whether it is chosen or required and how long it has been going on."),
    ("Pilot changes where overtime is heavy and not chosen. ",
     "Where overtime is heavy, sustained and not chosen, pilot changes in a few teams first, starting with demands "
     "that get in the way (red tape, unclear roles, understaffing) rather than work that stretches people; options "
     "include redistributing work, adding staff or adjusting schedules."),
    ("Track overtime and leaving together. ",
     "Report both side by side for each team, so leaders can compare teams that tried changes with similar teams "
     "that did not, before rolling changes out more widely."),
]
for i, (lead, body) in enumerate(recs, 1):
    p = P([(f"{i}.  ", "bd"), (lead, "bd"), (body, "")], after=4)
    p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.first_line_indent = Inches(-0.22)

# What Not to Conclude
H("What Not to Conclude: limits of a fictional sample")
nots = [
    ("It is a fictional sample. ", "IBM created these 1,470 employees for teaching. They describe no real workforce "
     "or employer."),
    ("It shows who left, not why. ", "Overtime may simply mark busy or understaffed roles, so cutting overtime alone "
     "may not reduce leaving."),
    ("Pay and career stage overlap. ", "Lower pay and junior roles go together, and the data can\u2019t say which "
     "matters."),
    ("It is not for decisions about individuals. ", "The patterns describe groups. Age, gender and marital status "
     "are descriptive only and must not be used for decisions about anyone."),
    ("Small groups are left out. ", "Groups of under 100 people were set aside as too small to act on."),
]
for lead, body in nots:
    P([(lead, "bd"), (body, "")], after=4)

# Footer
P([("Sources: ", "b"), ("Adams (1965); Greenhaus & Beutell (1985); Podsakoff, LePine & LePine (2007); Rubenstein, "
   "Eberly, Lee & Mitchell (2018). Numbers: analysis/exec/exec_findings_draft.md (QA-cleared Oct 9, 2026). "
   "Technical detail: analysis/exec/technical_appendix.md in the project repo.", "")],
  size=8.5, color=MUT, before=10, after=0)

cp = doc.core_properties
cp.title = "Overtime workers leave most: what IBM\u2019s fictional HR data suggests"
cp.author = "Iris"; cp.last_modified_by = "Iris"
import datetime as _dt
cp.created = cp.modified = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)
doc.save(OUT); print("saved", OUT)
