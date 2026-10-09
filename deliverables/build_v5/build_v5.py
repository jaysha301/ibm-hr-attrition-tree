"""Builds deliverables/attrition_narrative_v5.docx (v4 design + Ellis's research re-tie, research/exec_research_tiein.md; v1-v4 untouched).
Numbers ONLY from analysis/exec/exec_findings_draft.md (QA-cleared Oct 9, 2026).
Research paragraph, pay-context research sentence, four suggestions and footer sources use Ellis's wording
(exec_research_tiein.md sections 1, 2, 3, 5); pending Quinn's check of the research wording."""
import os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn

OUT = "../attrition_narrative_v5.docx"
if os.path.exists(OUT) and "--force" not in sys.argv:
    sys.exit(f"{OUT} exists; refusing to overwrite")

ACC = RGBColor(0xC2, 0x41, 0x0C); TXT = RGBColor(0x40, 0x40, 0x40)
MUT = RGBColor(0x6B, 0x6B, 0x6B); DARK = RGBColor(0x26, 0x26, 0x26)
doc = Document(); s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.left_margin = s.right_margin = Inches(0.75); s.top_margin = Inches(0.6); s.bottom_margin = Inches(0.55)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = TXT
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.space_before = Pt(0); st.paragraph_format.line_spacing = 1.1
W = 7.0

def P(parts, size=10.5, color=None, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    if isinstance(parts, str): parts = [(parts, "")]
    for t, f in parts:
        r = p.add_run(t); r.font.size = Pt(size)
        r.bold = "b" in f; r.font.color.rgb = ACC if "a" in f else (DARK if "d" in f else (color or TXT))
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p

def H(text, before=12):
    P([(text, "bd")], size=13, before=before, after=3, keep=True)

def img(path, width, alt, before=2, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    shape = p.add_run().add_picture(path, width=Inches(width))
    docpr = shape._inline.docPr; docpr.set("descr", alt); docpr.set("title", alt.split(":")[0])

# Title
P("IBM HR teaching data  \u00b7  fictional, simulated employees  \u00b7  executive summary", 9, MUT, after=3)
P([("Overtime workers", "ba"), (" leave most: what IBM\u2019s fictional HR data suggests", "bd")], size=21, after=10)

# Executive summary
P([("EXECUTIVE SUMMARY", "bd")], size=9.5, after=2)
P("In IBM\u2019s fictional teaching dataset of 1,470 simulated employees, 16% left, and overtime workers stand out: "
  "they are 28% of staff but 54% of everyone who left. If they left at the company average, overall attrition "
  "would fall from about 16% to about 12%, roughly 60 fewer leavers, as an illustration rather than a forecast. "
  "The first step is to find out where overtime is concentrated and why, then pilot changes in teams where "
  "overtime is heavy and not chosen. The key caution: these are simulated people, and the data show who left, not why.", after=6)
img("numbers_strip_v5.png", W, "Key figures: in IBM's fictional data, 16% of all employees left; 31% of overtime "
    "workers left, against 10% of staff without overtime.", before=4, after=4)

# What the Data Shows
H("What the Data Shows: overtime workers leave at about three times the rate of other staff")
P("In this fictional dataset, 416 people, 28% of staff, worked overtime. 31% of them left, about three times the "
  "rate of staff without overtime (10%). Together they account for 54% of everyone who left. It is the one large "
  "group that held up under every check the team ran.")
img("overtime_chart_v5.png", W, "Chart: in IBM's fictional data, overtime workers are 28% of all staff but 54% "
    "of everyone who left.", before=4, after=6)
P([("Context, not a separate finding: pay and career stage. ", "bd"),
   ("Lower-paid staff (roughly the bottom third, under about $3,000\u2013$4,000 a month) left at about 27%, against "
    "11% for everyone else, wherever the pay line is drawn in that range. But 95% of them are in the most junior job "
    "level, so this data cannot say whether pay or career stage matters. Most of their extra leaving is among those "
    "who also work overtime: without overtime, lower-paid staff leave at about the company average (16%), against 8% "
    "for better-paid staff without overtime. Research links both to leaving, since pay is a core reason to stay "
    "(March & Simon, 1958) and employees who are younger or newer to a company quit more often (Rubenstein et al., "
    "2018), so treat this as one picture, not two separate causes.", "")],
  size=9.5, color=MUT)

# Why It Matters and What to Do
H("Why It Matters and What to Do: overtime is the clearest place to start")
P("In this fictional data, the overtime group is large enough to move the overall number. If overtime workers "
  "left at the company average, overall attrition would fall from about 16% to about 12%, roughly 60 fewer leavers. That "
  "is an illustration of scale, not a forecast of what any change would achieve.")
P("Research on real workplaces gives reasons overtime could matter, and reasons for care. Long hours can drain "
  "time and energy from family life, a conflict linked to leaving (Greenhaus & Beutell, 1985; Rubenstein et al., "
  "2018). Equity theory holds that people who feel they give more than they get may leave to restore the balance "
  "(Adams, 1965). But research measures workload, not overtime, and finds it only weakly related to leaving, if "
  "anything slightly protective (Rubenstein et al., 2018). Demands like workload generally go with lower turnover; "
  "obstacles like red tape and unclear roles go with higher turnover (Podsakoff et al., 2007). This fictional data "
  "shows a far bigger overtime gap than research would predict: a pointer to where to look, not proof of cause or "
  "evidence about any real employer. This data does not record whether overtime was chosen or required, or how "
  "long it lasted, and in a real company those are the first things to find out.")
P([("Four suggestions to test, not proven fixes. ", "bd"),
   ("These are ideas to test, not interventions shown to work. Each applies to teams and to the overtime group as a "
    "whole, never to individuals and never selected by age, gender or marital status.", "")], after=5)
recs = [
    ("Review workload and staffing where overtime is concentrated. ",
     "Map overtime by team and role, and ask whether it is a predictable crunch or a chronic pattern, and whether it "
     "is chosen or required. Where it is heavy, sustained and not chosen, start with demands that get in the way "
     "(red tape, unclear roles, understaffing) rather than work that stretches people.",
     "Obstacles at work go with higher turnover while stretching demands do not (Podsakoff et al., 2007), and long "
     "hours can crowd out family life (Greenhaus & Beutell, 1985)."),
    ("Offer manager check-ins (\u201cstay conversations\u201d) to everyone in the overtime group. ",
     "Ask what keeps people and what might make them leave, then act on what is in the team\u2019s control. Offer "
     "them to the whole group, never to people picked by personal characteristics.",
     "Ties to colleagues, fit with the job and what people would give up by leaving help explain why people stay "
     "(Mitchell et al., 2001), and people\u2019s reasons for staying differ by job type (Hausknecht et al., 2009)."),
    ("Review pay and recognition for overtime effort, starting where overtime meets the lowest pay. ",
     "Check whether overtime is paid or recognized, and whether raises are handed out by a clear, fair process. "
     "Note that pay and career stage cannot be separated in this data.",
     "People stay while what they get outweighs what they give (March & Simon, 1958), and feeling under-rewarded "
     "may prompt people to leave (Adams, 1965)."),
    ("Pick one measure and pilot before rolling out. ",
     "Track the leaving rate of overtime workers by team, side by side with overtime levels. Try changes in a few "
     "teams first and compare them with similar teams that did not change, before rolling out more widely.",
     "Research finds most single drivers of leaving are modest (Griffeth et al., 2000; Rubenstein et al., 2018), "
     "and this data is fictional, so only a company\u2019s own tracked results can show whether a change helps."),
]
for i, (lead, body, why) in enumerate(recs, 1):
    p = P([(f"{i}.  ", "bd"), (lead, "bd"), (body, "")], after=1)
    p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.first_line_indent = Inches(-0.22)
    p = P([("Why: ", "b"), (why, "")], size=9, color=MUT, after=5)
    p.paragraph_format.left_indent = Inches(0.22)

# What Not to Conclude
H("What Not to Conclude: limits of a fictional sample")
nots = [
    ("It is a fictional sample. ", "IBM created these 1,470 employees for teaching. They describe no real workforce "
     "or employer."),
    ("It shows who left, not why. ", "Overtime may simply mark busy or understaffed roles, so cutting overtime alone "
     "may not reduce leaving."),
    ("Pay and career stage overlap. ", "Lower pay and junior roles go together, and the data can\u2019t say which "
     "matters."),
    ("It is not for decisions about individuals. ", "The patterns describe groups. Age, gender and marital status "
     "are descriptive only and must not be used for decisions about anyone."),
    ("Small groups are left out. ", "Groups of under 100 people were set aside as too small to act on."),
]
for lead, body in nots:
    P([(lead, "bd"), (body, "")], after=4)

# Footer (Ellis's short-form sources line, exec_research_tiein.md line 56)
P([("Sources: ", "b"), ("Adams (1965); Greenhaus & Beutell (1985); Griffeth, Hom & Gaertner (2000); Hausknecht, "
   "Rodda & Howard (2009); March & Simon (1958); Mitchell, Holtom, Lee, Sablynski & Erez (2001); Podsakoff, LePine & "
   "LePine (2007); Rubenstein, Eberly, Lee & Mitchell (2018). Numbers: analysis/exec/exec_findings_draft.md "
   "(QA-cleared Oct 9, 2026). Technical detail: analysis/exec/technical_appendix.md in the project repo.", "")],
  size=8.5, color=MUT, before=10, after=0)

cp = doc.core_properties
cp.title = "Overtime workers leave most: what IBM\u2019s fictional HR data suggests"
cp.author = "Iris"; cp.last_modified_by = "Iris"
import datetime as _dt
cp.created = cp.modified = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)
doc.save(OUT); print("saved", OUT)
