"""Builds deliverables/attrition_narrative_v7.docx (exec2 rebuild; v1-v6 untouched).
Numbers ONLY from analysis/exec2/exec_findings_v2_draft.md (QA-cleared by Quinn, Oct 9, 2026).
Research paragraph and four suggestions reused verbatim from v6 (Quinn-cleared). The two
'Research finds that both lower pay...' sentences (subgroup context, watching box) are PROVISIONAL
placeholders pending Ellis's research for the subgroup and the career-stage box."""
import os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn

OUT = "../attrition_narrative_v7.docx"
if os.path.exists(OUT) and "--force" not in sys.argv:
    sys.exit(f"{OUT} exists; refusing to overwrite")

ACC = RGBColor(0xC2, 0x41, 0x0C); TXT = RGBColor(0x40, 0x40, 0x40)
MUT = RGBColor(0x6B, 0x6B, 0x6B); DARK = RGBColor(0x26, 0x26, 0x26)
doc = Document(); s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.left_margin = s.right_margin = Inches(0.75); s.top_margin = Inches(0.5); s.bottom_margin = Inches(0.5)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10); st.font.color.rgb = TXT
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.space_before = Pt(0); st.paragraph_format.line_spacing = 1.05
W = 7.0

def P(parts, size=10, color=None, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    if isinstance(parts, str): parts = [(parts, "")]
    for t, f in parts:
        r = p.add_run(t); r.font.size = Pt(size)
        r.bold = "b" in f; r.font.color.rgb = ACC if "a" in f else (DARK if "d" in f else (color or TXT))
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p

def H(text, before=9):
    P([(text, "bd")], size=13, before=before, after=3, keep=True)

def img(path, width, alt, before=2, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    shape = p.add_run().add_picture(path, width=Inches(width))
    docpr = shape._inline.docPr; docpr.set("descr", alt); docpr.set("title", alt.split(":")[0])

from docx.oxml import OxmlElement
ACC2 = RGBColor(0x7C, 0x2D, 0x12)

def shaded_box(parts_list, fill="F2F2F2"):
    t = doc.add_table(rows=1, cols=1); t.autofit = False
    c = t.rows[0].cells[0]; c.width = Inches(W)
    pr = c._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); pr.append(shd)
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", 110), ("bottom", 90), ("start", 160), ("end", 160)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    pr.append(mar)
    tb = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:val"), "nil"); b.append(e)
    tb.append(b)
    first = True
    for parts, size, after in parts_list:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        for txt, f in parts:
            r = p.add_run(txt); r.font.size = Pt(size); r.bold = "b" in f
            r.font.color.rgb = DARK if "d" in f else MUT
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

PLACEHOLDER = ("Research finds that both lower pay and being newer to a company go with more leaving "
               "(Rubenstein et al., 2018), so treat this as one picture, not two separate explanations.")  # PROVISIONAL

# Title
P("IBM HR teaching data  \u00b7  fictional, simulated employees  \u00b7  executive summary", 9, MUT, after=3)
P([("Overtime workers", "ba"), (" leave most in IBM\u2019s fictional HR data, lower-paid junior ones most of all", "bd")],
  size=20, after=10)

# Executive summary
P([("EXECUTIVE SUMMARY", "bd")], size=9.5, after=2)
P("In IBM\u2019s fictional teaching dataset of 1,470 simulated employees, 16% left. Overtime workers are 28% of staff "
  "but 54% of everyone who left, and within that group lower-paid, mostly junior staff leave most: about 55% of them "
  "left, close to 30% of all leavers. As an illustration, not a forecast, if overtime workers left at the company "
  "average, overall attrition would fall from about 16% to about 12%, roughly 60 fewer leavers. The first step is to "
  "find out where overtime is concentrated and why, then pilot changes in teams where overtime is heavy and not "
  "chosen; the key caution is that these are simulated people, and the data show who left, not why.", after=6)

# What the Data Shows
H("What the Data Shows: overtime workers leave most, lower-paid junior ones most of all")
P("In this fictional dataset, 416 people, 28% of staff, worked overtime. 31% of them left, about 1.9 times the "
  "company rate of 16% and about 3 times the rate of staff without overtime (10%). Together they account for 54% of "
  "everyone who left.")
P("Within overtime, lower-paid, mostly junior staff leave most. Overtime workers earning under roughly $3,000 to "
  "$3,500 a month (about 115 to 130 people, 8% to 9% of staff) left at about 55%, about 2.7 to 2.9 times the rate of "
  "the other overtime workers (about 20%). They account for close to 30% of everyone who left. They are part of the "
  "overtime finding, not a separate pattern.")
img("leaving_chart_v7.png", 6.6, "Chart: in IBM's fictional data, 16% of all staff and 10% of staff without overtime "
    "left, against 31% of the 416 overtime workers and about 55% of the about 115 to 130 lower-paid, junior overtime "
    "workers. Overtime workers are 54% of everyone who left; the lower-paid, junior ones, inside that group, are close "
    "to 30%.", before=4, after=6)
P([("Context: one career-stage picture. ", "bd"),
   ("Pay, job level, tenure and total experience move together, and the data cannot say which of them matters. "
    "Almost all of these lower-paid overtime workers (97%) are in the most junior job level, so pay and job level "
    "cannot be separated: overtime workers in the most junior job level (156 people) left at 53%, against 17% for "
    "the 260 overtime workers at higher job levels, close to the company average of 16%. " + PLACEHOLDER, "")],
  size=9.5, color=MUT)
shaded_box([
    ([("Worth watching, not a cleared finding: newer staff", "bd")], 10, 3),
    ([("Staff with 1 year or less at the company (215 people, 15% of staff) left at 35%, about 40 people above the "
       "company rate; about a third of them (32%) work overtime. The group is large, but the exact group isn\u2019t "
       "stable across repeated analyses, because new hires, short tenure, junior level and low pay overlap. Treat them "
       "as one career-stage picture to watch, not a proven cause; it is not added to the overtime figures. Among staff "
       "without overtime, newer staff also leave more often: a real difference inside that group, but modest against "
       "the company rate. " + PLACEHOLDER, "")], 9.5, 0),
])

# Why It Matters and What to Do
H("Why It Matters and What to Do: overtime is the clearest place to start")
P("In this fictional data, the overtime group is large enough to move the overall number. If overtime workers left at "
  "the company average, overall attrition would fall from about 16% to about 12%, roughly 60 fewer leavers. On the "
  "same illustration, the lower-paid, mostly junior overtime workers would take it from about 16% to about 13%, "
  "roughly 45 to 50 fewer leavers. They are already inside the overtime group, so the two illustrations overlap and "
  "are not added. Both are illustrations of scale, not forecasts of what any change would achieve.")
P("Research on real workplaces gives reasons overtime could matter, and reasons for care. Long hours can drain "
  "time and energy from family life (Greenhaus & Beutell, 1985), and that kind of conflict is linked to leaving "
  "(Rubenstein et al., 2018). Equity theory holds that people who feel they give more than they get may leave to restore the balance "
  "(Adams, 1965); this data does not record how people felt. But research measures workload, not overtime, and finds it only weakly related to leaving, if "
  "anything slightly protective (Rubenstein et al., 2018). Demands like workload generally go with lower turnover; "
  "obstacles like red tape and unclear roles go with higher turnover (Podsakoff et al., 2007). Research on "
  "workload would not lead us to expect a gap this large, so treat this fictional result as a pointer to where to "
  "look, not proof of cause or evidence about any real employer. This data does not record whether overtime was chosen or required, or how "
  "long it lasted, and in a real company those are the first things to find out.")
P([("Four suggestions to test, not proven fixes. ", "bd"),
   ("Each applies to teams and to the overtime group as a "
    "whole, never to individuals and never selected by age, gender or marital status.", "")], after=5)
recs = [
    ("Review workload and staffing where overtime is concentrated. ",
     "Map overtime by team and role, and ask whether it is a predictable crunch or a chronic pattern, and whether it "
     "is chosen or required. Where it is heavy, sustained and not chosen, start with demands that get in the way "
     "(red tape, unclear roles, understaffing) rather than work that stretches people.",
     "Obstacles at work go with higher turnover while stretching demands do not (Podsakoff et al., 2007), and long "
     "hours can crowd out family life (Greenhaus & Beutell, 1985)."),
    ("Offer manager check-ins (\u201cstay conversations\u201d) to everyone in the overtime group. ",
     "Ask what keeps people and what might make them leave, then act on what is in the team\u2019s control. Offer "
     "them to the whole group, never to people picked by personal characteristics.",
     "Ties to colleagues, fit with the job and what people would give up by leaving help explain why people stay "
     "(Mitchell et al., 2001), and people\u2019s reasons for staying differ by job type (Hausknecht et al., 2009)."),
    ("Review pay and recognition for overtime effort across the overtime group. ",
     "Check whether overtime is paid or recognized, and whether raises are handed out by a clear, fair process. "
     "Note that pay and career stage cannot be separated in this data.",
     "Classic theory holds that people stay while what they get outweighs what they give (March & Simon, 1958), "
     "and that feeling under-rewarded can prompt people to leave (Adams, 1965)."),
    ("Pick one measure and pilot before rolling out. ",
     "Track the leaving rate of overtime workers by team, side by side with overtime levels. Try changes in a few "
     "teams first and compare them with similar teams that did not change, before rolling out more widely.",
     "Research finds that most single factors are only modestly related to leaving (Griffeth et al., 2000; Rubenstein et al., 2018), "
     "and this data is fictional, so only a company\u2019s own tracked results can show whether a change helps."),
]
for i, (lead, body, why) in enumerate(recs, 1):
    p = P([(f"{i}.  ", "bd"), (lead, "bd"), (body, "")], after=1)
    p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.first_line_indent = Inches(-0.22)
    p = P([("Why: ", "b"), (why, "")], size=9, color=MUT, after=4)
    p.paragraph_format.left_indent = Inches(0.22)

# What Not to Conclude
H("What Not to Conclude: limits of a fictional sample")
nots = [
    ("It is a fictional sample. ", "IBM created these 1,470 employees for teaching. They describe no real workforce "
     "or employer."),
    ("It shows who left, not why. ", "Overtime may simply mark busy or understaffed roles, and pay moves with job "
     "level and experience."),
    ("Pay and career stage overlap. ", "Lower pay, junior job level, short tenure and little experience go together, "
     "and the data can\u2019t say which matters."),
    ("It is not for decisions about individuals. ", "The patterns describe groups. Age, gender and marital status "
     "are descriptive only and never grounds for targeting anyone."),
    ("Small groups are left out. ", "Groups under 100 people were set aside as too small to act on."),
]
for lead, body in nots:
    P([(lead, "bd"), (body, "")], after=3)

# Footer (v6 sources; all eight still cited via the reused research paragraph and suggestions)
P([("Sources: ", "b"), ("Adams (1965); Greenhaus & Beutell (1985); Griffeth, Hom & Gaertner (2000); Hausknecht, "
   "Rodda & Howard (2009); March & Simon (1958); Mitchell, Holtom, Lee, Sablynski & Erez (2001); Podsakoff, LePine & "
   "LePine (2007); Rubenstein, Eberly, Lee & Mitchell (2018). Numbers: analysis/exec2/exec_findings_v2_draft.md "
   "(QA-cleared Oct 9, 2026). Technical detail: analysis/exec2/technical_appendix.md in the project repo.", "")],
  size=8.5, color=MUT, before=6, after=0)

cp = doc.core_properties
cp.title = "Overtime workers leave most in IBM\u2019s fictional HR data, lower-paid junior ones most of all"
cp.author = "Iris"; cp.last_modified_by = "Iris"
import datetime as _dt
cp.created = cp.modified = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)
doc.save(OUT); print("saved", OUT)
